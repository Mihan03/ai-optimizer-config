"""Собирает khl_simulation_v2.xlsx из исходной книги v1.

python scripts/build_workbook.py --src khl_simulation.xlsx --out khl_simulation_v2.xlsx
"""
from __future__ import annotations

import argparse
import sys
from dataclasses import asdict, fields
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from khl_model import MatchModel, Params  # noqa: E402
from khl_model.betting import (POOLS, Offer, describe, evaluate_offer, passes, pnl,  # noqa: E402
                               select_bets, settle, stake_for)

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN, RED, GREY, YELLOW = "C6EFCE", "FFC7CE", "D9D9D9", "FFEB9C"
PCT, ODDS, MONEY = "0.0%", "0.00", "#,##0"

RESULTS = ["ЗАШЛО", "НЕ ЗАШЛО", "ВОЗВРАТ", "ОЖИДАНИЕ", "ОТМЕНА"]
FIRST, LAST = 8, 5000  # строки журнала

JOURNAL_COLS = [
    ("ID", 22), ("Дата", 11), ("Время МСК", 9), ("Хозяева", 16), ("Гости", 16), ("Cutoff (UTC+5)", 17),
    ("Версия", 7), ("Пул", 10), ("Рынок", 9), ("Выбор", 7), ("Линия", 7), ("Ставка", 28),
    ("Коэффициент", 8), ("Букмекер", 10), ("Время кэфа", 17), ("λ хозяев", 8), ("λ гостей", 8),
    ("p модели", 8), ("q рынка", 8), ("p итог", 8), ("p возврата", 8), ("Fair кэф", 8), ("EV", 8),
    ("EV худший", 8), ("Kelly", 7), ("Банк пула до ставки", 10), ("Ставка, у.е.", 9),
    ("Кэф закрытия", 8), ("q закрытия", 8), ("CLV", 8), ("Счёт 60м", 8), ("Счёт итог", 10),
    ("Результат", 12), ("P&L", 9), ("Обоснование (из признаков модели)", 60),
]
C = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(JOURNAL_COLS)}


def rng(col: str) -> str:
    return f"${C[col]}${FIRST}:${C[col]}${LAST}"


def style_header(ws, row: int, ncols: int, fill=HEAD_FILL, font=HEAD_FONT) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = fill, font, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def add_result_colors(ws, ref: str, first_cell: str) -> None:
    for text, color in (("ЗАШЛО", GREEN), ("НЕ ЗАШЛО", RED), ("ВОЗВРАТ", GREY), ("ОТМЕНА", GREY), ("ОЖИДАНИЕ", YELLOW)):
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'{first_cell}="{text}"'], fill=PatternFill("solid", fgColor=color)))


# ---------------------------------------------------------------- журнал v2
def build_journal(wb) -> None:
    ws = wb.create_sheet("KHL", 0)
    summary = ["Пул", "Начальный банк", "Текущий банк", "Ставок", "Рассчитано", "Зашло", "Не зашло",
               "Возврат", "Ожидание", "Оборот", "P&L", "Yield", "Средний кэф", "Средний EV (ожид.)",
               "Средний CLV", "Brier модели", "Brier рынка", "Hit rate"]
    for i, h in enumerate(summary, 1):
        ws.cell(row=1, column=i, value=h)
    style_header(ws, 1, len(summary))

    H, R, S, M, T, Q = rng("Пул"), rng("Результат"), rng("Ставка, у.е."), rng("Коэффициент"), rng("p итог"), rng("q рынка")
    EV, QC = rng("EV"), rng("q закрытия")
    for r, pool in enumerate(list(POOLS) + ["ИТОГО"], start=2):
        total = pool == "ИТОГО"
        cond = f'({H}<>"")' if total else f"({H}=$A{r})"
        crit = '"<>"' if total else f"$A{r}"
        ws.cell(row=r, column=1, value=pool)
        f = {
            "B": "=SUM(B2:B4)" if total else "=1000",
            "C": f"=B{r}+K{r}",
            "D": f"=COUNTIFS({H},{crit})",
            "E": f"=F{r}+G{r}+H{r}",
            "F": f'=COUNTIFS({H},{crit},{R},"ЗАШЛО")',
            "G": f'=COUNTIFS({H},{crit},{R},"НЕ ЗАШЛО")',
            "H": f'=COUNTIFS({H},{crit},{R},"ВОЗВРАТ")',
            "I": f'=COUNTIFS({H},{crit},{R},"ОЖИДАНИЕ")',
            "J": f'=SUMIFS({S},{H},{crit},{R},"ЗАШЛО")+SUMIFS({S},{H},{crit},{R},"НЕ ЗАШЛО")+SUMIFS({S},{H},{crit},{R},"ВОЗВРАТ")',
            "K": f'=SUMPRODUCT({cond}*({R}="ЗАШЛО")*{S}*({M}-1))-SUMIFS({S},{H},{crit},{R},"НЕ ЗАШЛО")',
            "L": f'=IF(J{r}>0,K{r}/J{r},"")',
            "M": f'=IFERROR(AVERAGEIFS({M},{H},{crit}),"")',
            "N": f'=IFERROR(AVERAGEIFS({EV},{H},{crit}),"")',
            "O": f'=IFERROR(SUMPRODUCT({cond}*({QC}<>"")*({M}*{QC}-1))/COUNTIFS({H},{crit},{QC},"<>"),"")',
            "P": f'=IFERROR(SUMPRODUCT({cond}*(({R}="ЗАШЛО")+({R}="НЕ ЗАШЛО"))*({T}-({R}="ЗАШЛО"))^2)/(F{r}+G{r}),"")',
            "Q": f'=IFERROR(SUMPRODUCT({cond}*(({R}="ЗАШЛО")+({R}="НЕ ЗАШЛО"))*({Q}-({R}="ЗАШЛО"))^2)/(F{r}+G{r}),"")',
            "R": f'=IF(F{r}+G{r}>0,F{r}/(F{r}+G{r}),"")',
        }
        for col, formula in f.items():
            ws[f"{col}{r}"] = formula
        for col in "BCJK":
            ws[f"{col}{r}"].number_format = MONEY
        for col in "LNOR":
            ws[f"{col}{r}"].number_format = PCT
        ws[f"M{r}"].number_format = ODDS
        for col in "PQ":
            ws[f"{col}{r}"].number_format = "0.000"
        for c in range(1, len(summary) + 1):
            ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=1).font = BOLD
        if total:
            for c in range(1, len(summary) + 1):
                ws.cell(row=r, column=c).fill = SUB_FILL
                ws.cell(row=r, column=c).font = BOLD
    ws.conditional_formatting.add("K2:K5", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add("K2:K5", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add("C2:C4", CellIsRule(operator="lessThan", formula=["B2*0.5"], fill=PatternFill("solid", fgColor=RED)))

    for i, (name, width) in enumerate(JOURNAL_COLS, 1):
        ws.cell(row=7, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 10 if i <= len(summary) else width)
    style_header(ws, 7, len(JOURNAL_COLS))
    ws.row_dimensions[7].height = 32
    ws.freeze_panes = "F8"

    def col_ref(name):
        return f"{C[name]}{FIRST}:{C[name]}{LAST}"

    for name, fmt in (("p модели", PCT), ("q рынка", PCT), ("p итог", PCT), ("p возврата", PCT), ("EV", PCT),
                      ("EV худший", PCT), ("Kelly", PCT), ("q закрытия", PCT), ("CLV", PCT),
                      ("Коэффициент", ODDS), ("Fair кэф", ODDS), ("Кэф закрытия", ODDS),
                      ("λ хозяев", ODDS), ("λ гостей", ODDS), ("Ставка, у.е.", MONEY), ("P&L", MONEY),
                      ("Банк пула до ставки", MONEY)):
        for row in range(FIRST, FIRST + 300):
            ws[f"{C[name]}{row}"].number_format = fmt

    add_result_colors(ws, col_ref("Результат"), f"{C['Результат']}{FIRST}")
    ws.conditional_formatting.add(col_ref("CLV"), CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(col_ref("CLV"), CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))

    for name, options in (("Пул", POOLS), ("Результат", RESULTS),
                          ("Рынок", ("1X2", "DC", "TOTAL", "HCP", "TEAM_TOTAL", "ML_OT"))):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(col_ref(name))


# ---------------------------------------------------------------- прогнозы по всем матчам
FC_COLS = [("ID", 22), ("Дата", 11), ("Хозяева", 16), ("Гости", 16), ("Источник λ", 14), ("λ хозяев", 8),
           ("λ гостей", 8), ("P1 60м", 8), ("X 60м", 8), ("P2 60м", 8), ("П1 с ОТ/Б", 9), ("П2 с ОТ/Б", 9),
           ("ТМ 5.5", 8), ("q1 рынка", 8), ("qX рынка", 8), ("q2 рынка", 8), ("Счёт 60м", 9), ("Исход 60м", 9),
           ("Топ-счета", 40), ("80% интервал тотала", 10)]


def build_forecasts(wb, bets: pd.DataFrame, params: Params) -> None:
    ws = wb.create_sheet("Прогнозы", 1)
    labels = ["Матчей с исходом", "Brier модели (1X2)", "Brier рынка (1X2)", "Log loss модели", "Log loss рынка",
              "Средняя P(X) модели", "Средняя qX рынка", "Фактическая доля ничьих"]
    a, b = 6, 3000
    P1, PX, P2 = (f"$H${a}:$H${b}", f"$I${a}:$I${b}", f"$J${a}:$J${b}")
    Q1, QX, Q2 = (f"$N${a}:$N${b}", f"$O${a}:$O${b}", f"$P${a}:$P${b}")
    Y = f"$R${a}:$R${b}"
    has_q = f'({Y}<>"")*({Q1}<>"")'
    formulas = [
        f'=COUNTA({Y})',
        f'=IFERROR(SUMPRODUCT(({Y}<>"")*(({P1}-({Y}="П1"))^2+({PX}-({Y}="X"))^2+({P2}-({Y}="П2"))^2))/A2,"")',
        f'=IFERROR(SUMPRODUCT({has_q}*(({Q1}-({Y}="П1"))^2+({QX}-({Y}="X"))^2+({Q2}-({Y}="П2"))^2))/SUMPRODUCT({has_q}),"")',
        f'=IFERROR(-SUMPRODUCT(({Y}="П1")*LN({P1}+({Y}<>"П1")+({P1}=""))+({Y}="X")*LN({PX}+({Y}<>"X")+({PX}=""))+({Y}="П2")*LN({P2}+({Y}<>"П2")+({P2}="")))/A2,"")',
        f'=IFERROR(-SUMPRODUCT({has_q}*(({Y}="П1")*LN({Q1}+({Y}<>"П1")+({Q1}=""))+({Y}="X")*LN({QX}+({Y}<>"X")+({QX}=""))+({Y}="П2")*LN({Q2}+({Y}<>"П2")+({Q2}=""))))/SUMPRODUCT({has_q}),"")',
        f'=IFERROR(AVERAGE({PX}),"")',
        f'=IFERROR(AVERAGE({QX}),"")',
        f'=IFERROR(COUNTIF({Y},"X")/A2,"")',
    ]
    for i, (lab, f) in enumerate(zip(labels, formulas), 1):
        ws.cell(row=1, column=i, value=lab)
        ws.cell(row=2, column=i, value=f).border = BORDER
        ws.cell(row=2, column=i).number_format = "0" if i == 1 else ("0.000" if i <= 5 else PCT)
    style_header(ws, 1, len(labels))
    ws.row_dimensions[1].height = 32
    ws["A3"] = ("Главная проверка качества — по ВСЕМ матчам, а не только по ставкам: модель полезна, только если её "
                "Brier/log loss не хуже рыночного, а средняя P(X) близка к фактической доле ничьих. "
                "Исход писать текстом: П1 / X / П2.")
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:H3")
    ws.row_dimensions[3].height = 45

    for i, (name, width) in enumerate(FC_COLS, 1):
        ws.cell(row=5, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_header(ws, 5, len(FC_COLS))
    ws.row_dimensions[5].height = 32
    ws.freeze_panes = "E6"

    matches = bets.drop_duplicates("match_id")
    for r, (_, m) in enumerate(matches.iterrows(), start=a):
        mm = MatchModel(m.lam_home, m.lam_away, params)
        p1, px, p2 = mm.p_1x2()
        w1, w2 = mm.p_winner()
        h, g = (int(x) for x in m.score_60.split(":"))
        lo, hi = mm.total_interval()
        values = [m.match_id, m.match_id.split("-")[3] + "." + m.match_id.split("-")[2] + "." + m.match_id.split("-")[1],
                  m.home, m.away, "v1 (ретро)", m.lam_home, m.lam_away, p1, px, p2, w1, w2,
                  mm.market_prob("TOTAL", "U", 5.5)[0], None, None, None, m.score_60,
                  "П1" if h > g else ("X" if h == g else "П2"),
                  ", ".join(f"{s} ({p:.1%})" for s, p in mm.top_scores(5)), f"[{lo}, {hi}]"]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if 8 <= c <= 16:
                cell.number_format = PCT
            if c in (6, 7):
                cell.number_format = ODDS


# ---------------------------------------------------------------- пересчёт v1 → v2
def build_retro(wb, bets: pd.DataFrame, params: Params) -> dict:
    ws = wb.create_sheet("Пересчёт v1→v2", 2)
    ws["A1"] = "Пересчёт 21 ставки v1.0 по правилам v2.0"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("λ взяты из v1 (в v2 они считаются по рейтингам команд). q рынка: для главных событий — из журнала v1 "
                "(без маржи), для надёжных и смелых — оценка по одному коэффициенту при марже 5% (вторая сторона "
                "рынка в v1 не сохранялась). Выборка из 7 матчей ничего не доказывает — пересчёт показывает, как "
                "меняется логика отбора.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:T2")
    ws.row_dimensions[2].height = 48

    cols = [("Матч", 26), ("Пул v1", 10), ("Ставка", 30), ("Кэф", 7), ("Результат", 11), ("p v1", 8),
            ("EV v1 (по модели)", 9), ("p модели v2", 8), ("q рынка", 8), ("Источник q", 14), ("p итог v2", 8),
            ("EV v2", 8), ("EV худший", 8), ("Пул v2", 10), ("Решение v2", 13), ("Причина", 38),
            ("P&L v1 (1000)", 10), ("Ставка v2", 9), ("P&L v2", 8)]
    for i, (name, width) in enumerate(cols, 1):
        ws.cell(row=4, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = width
    style_header(ws, 4, len(cols))
    ws.row_dimensions[4].height = 32

    evals = []
    for _, b in bets.iterrows():
        line = None if pd.isna(b.line) else float(b.line)
        offer = Offer(b.match_id, b.market, str(b.selection), line, float(b.odds),
                      q_market=None if pd.isna(b.q_market) else float(b.q_market))
        evals.append(evaluate_offer(b.lam_home, b.lam_away, offer, params))
    chosen = {id(e) for e in select_bets(evals, params)}

    totals = {"v1": 0.0, "v2": 0.0, "v2_bets": 0, "v1_by_pool": {}}
    for r, (b, e) in enumerate(zip(bets.itertuples(), evals), start=5):
        line = None if pd.isna(b.line) else float(b.line)
        res = settle(b.market, str(b.selection), line, b.score_60, b.score_final)
        res_ru = {"WIN": "ЗАШЛО", "LOSS": "НЕ ЗАШЛО", "PUSH": "ВОЗВРАТ"}[res]
        p1 = pnl(res, 1000, b.odds)
        totals["v1"] += p1
        totals["v1_by_pool"][b.pool_v1] = totals["v1_by_pool"].get(b.pool_v1, 0) + p1
        if id(e) in chosen:
            decision, reason = "СТАВКА", "EV ≥ порога и устойчив к стресс-сценариям"
        elif e.ev < params.ev_min:
            decision = "ПРОПУСК"
            reason = (f"EV {e.ev:+.1%} < {params.ev_min:.0%}" +
                      ("; по самой модели v1 EV тоже < 0" if b.p_v1 * b.odds - 1 < 0 else ""))
        elif not passes(e, params):
            decision, reason = "ПРОПУСК", f"неустойчиво: в стресс-сценарии EV {e.ev_worst:+.1%}"
        else:
            decision, reason = "ПРОПУСК", "на матч уже выбрана ставка с бо́льшим EV"
        stake = stake_for(e, params.start_bank, params) if decision == "СТАВКА" else 0
        p2 = pnl(res, stake, b.odds) if stake else 0
        totals["v2"] += p2
        totals["v2_bets"] += bool(stake)
        row = [f"{b.home} – {b.away}", b.pool_v1, describe(b.market, str(b.selection), line, b.home, b.away),
               b.odds, res_ru, b.p_v1, b.p_v1 * b.odds - 1, e.p_model[0], e.q_market,
               "журнал v1" if not pd.isna(b.q_market) else "оценка, маржа 5%", e.p_final[0], e.ev, e.ev_worst,
               e.pool, decision, reason, p1, stake, p2]
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c in (6, 7, 8, 9, 11, 12, 13):
                cell.number_format = PCT
            if c in (17, 18, 19):
                cell.number_format = MONEY
    end = 4 + len(bets)
    add_result_colors(ws, f"E5:E{end}", "E5")
    for col in ("G", "L", "M"):
        ws.conditional_formatting.add(f"{col}5:{col}{end}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))
        ws.conditional_formatting.add(f"{col}5:{col}{end}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=Font(color="006100")))
    ws.freeze_panes = "B5"

    # Ничьи: модель v1, v2 и рынок
    r0 = end + 3
    ws.cell(row=r0 - 1, column=1, value="Вероятность ничьей за 60 минут: v1 vs v2 vs рынок").font = Font(bold=True, size=12)
    hdr = ["Матч", "X v1 (чистый Пуассон)", "X v2", "X рынка ≈ 1/кэф(1X) − 1/кэф(П1)", "Комментарий"]
    for i, h in enumerate(hdr, 1):
        ws.cell(row=r0, column=i, value=h)
    style_header(ws, r0, len(hdr))
    plain = params.replace(draw_inflation=0.0, empty_net_shift=0.0)
    r = r0 + 1
    for mid, g in bets.groupby("match_id", sort=False):
        one = g[(g.market == "1X2") & (g.selection.astype(str) == "1")]
        dc = g[(g.market == "DC") & (g.selection == "1X")]
        if one.empty or dc.empty:
            continue
        m = g.iloc[0]
        x1 = MatchModel(m.lam_home, m.lam_away, plain).p_1x2()[1]
        x2 = MatchModel(m.lam_home, m.lam_away, params).p_1x2()[1]
        xm = 1 / dc.iloc[0].odds - 1 / one.iloc[0].odds
        for c, v in enumerate([f"{m.home} – {m.away}", x1, x2, xm,
                               "маржа в обеих ценах примерно сокращается"], 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c in (2, 3, 4):
                cell.number_format = PCT
        r += 1

    # Итоги
    r += 2
    ws.cell(row=r, column=1, value="Итоги").font = Font(bold=True, size=12)
    lines = [
        ("P&L v1, все 21 ставка по 1000", totals["v1"]),
        *[(f"   в т.ч. {k}", v) for k, v in totals["v1_by_pool"].items()],
        ("Ставок v2", totals["v2_bets"]),
        ("P&L v2", totals["v2"]),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ws.cell(row=r + i, column=1, value=k)
        ws.cell(row=r + i, column=2, value=v).number_format = MONEY
    ws.cell(row=r + len(lines) + 2, column=1, value=(
        "Вывод: после поправки на ничьи и смешивания с рынком ни одна ставка v1 не проходит порог EV ≥ 3% с "
        "проверкой устойчивости. Все 7 «надёжных» ставок имели EV < 0 даже по вероятностям самой v1. "
        "Результат +3010 у «смелых» — дисперсия на 7 ставках, а не перевес."))
    ws.cell(row=r + len(lines) + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + len(lines) + 2, start_column=1, end_row=r + len(lines) + 2, end_column=12)
    ws.row_dimensions[r + len(lines) + 2].height = 48
    return totals


# ---------------------------------------------------------------- рейтинги, настройки, промпт, архив
def build_ratings(wb) -> None:
    ws = wb.create_sheet("Рейтинги", 3)
    ws["A1"] = "Рейтинги команд на дату прогноза (python -m khl_model fit … ; копировать вывод целиком, руками не править)"
    ws["A1"].font = BOLD
    ws["A2"] = "Дата среза:"
    ws["C2"] = "μ:"
    ws["E2"] = "h (дом):"
    cols = [("Команда", 22), ("att", 9), ("def", 9), ("Голы за 60м vs средний", 12), ("Пропущено за 60м vs средний", 12),
            ("Взвешенных матчей", 11)]
    for i, (name, w) in enumerate(cols, 1):
        ws.cell(row=4, column=i, value=name)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, 4, len(cols))
    ws.row_dimensions[4].height = 32


SETTINGS_DESC = {
    "mu": ("ФИТ", "Базовый log-λ команды за 60 минут; фитится вместе с рейтингами"),
    "home_adv": ("ФИТ", "log-преимущество своей площадки; фитится вместе с рейтингами"),
    "ridge": ("ПРАВИЛО", "L2-регуляризация att/def — сжимает рейтинги к среднему лиги"),
    "half_life_days": ("ПРАВИЛО", "Период полураспада веса матча, дней"),
    "season_carryover": ("ПРАВИЛО", "Вес матчей прошлых сезонов (рейтинги в начале сезона регрессируют к среднему)"),
    "draw_inflation": ("PROVISIONAL", "θ: множитель (1+θ) на ничейные счета. 0.40 подобрано так, чтобы P(X) совпадала с рыночной ~20–23%; калибровать по истории"),
    "empty_net_shift": ("PROVISIONAL", "Доля побед в 1 шайбу, превращающихся в +2 из-за пустых ворот; калибровать (нужно ≥2 сезонов)"),
    "max_goals": ("ПРАВИЛО", "Размер матрицы счетов"),
    "ot_goal_prob": ("PROVISIONAL", "P(решающий гол в ОТ 3×3 | ничья за 60 минут); иначе буллиты"),
    "ot_strength_shrink": ("ПРАВИЛО", "Насколько сила команд переносится на ОТ (0 — монетка, 1 — как в основное время)"),
    "shootout_home": ("PROVISIONAL", "P(хозяева выигрывают серию буллитов)"),
    "backup_goalie_effect": ("PROVISIONAL", "+к λ соперника, если в воротах запасной вратарь"),
    "b2b_attack_effect": ("PROVISIONAL", "Второй матч за 2 дня: изменение своей λ"),
    "b2b_defense_effect": ("PROVISIONAL", "Второй матч за 2 дня: изменение λ соперника"),
    "travel_tz_effect": ("PROVISIONAL", "Изменение своей λ за каждый час смены часового пояса (до 4 ч)"),
    "market_weight_model": ("PROVISIONAL", "Вес модели при смешивании с рынком (0 — только рынок, 1 — только модель). Калибровать по log loss"),
    "default_overround": ("ПРАВИЛО", "Маржа для оценки q, если известна только одна сторона рынка"),
    "ev_min": ("ПРАВИЛО", "Минимальный EV после смешивания с рынком"),
    "require_robust": ("ПРАВИЛО", "Требовать EV ≥ 0 во всех стресс-сценариях (λ ±7%, θ ×0.5/×1.5, вес рынка +0.2)"),
    "max_bets_per_match": ("ПРАВИЛО", "Ставки на один матч коррелированы — не больше одной"),
    "reliable_max_odds": ("ПРАВИЛО", "Пул «Надёжное»: кэф ниже этого значения"),
    "bold_min_odds": ("ПРАВИЛО", "Пул «Смелое»: кэф не ниже этого значения; между — «Главное»"),
    "start_bank": ("ПРАВИЛО", "Стартовый банк каждого пула"),
    "kelly_fraction": ("ПРАВИЛО", "Доля Келли"),
    "max_stake_frac": ("ПРАВИЛО", "Потолок ставки — доля ТЕКУЩЕГО банка пула"),
    "min_stake": ("ПРАВИЛО", "Ставки меньше — не делаются"),
}


def build_settings(wb, params: Params) -> None:
    ws = wb.create_sheet("Настройки")
    for i, (h, w) in enumerate((("Параметр", 24), ("Значение", 14), ("Статус", 13), ("Описание", 95)), 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header(ws, 1, 4)
    r = 2
    for f in fields(Params):
        if f.name in ("notes",):
            continue
        status, desc = SETTINGS_DESC.get(f.name, ("", ""))
        if f.name == "version":
            status, desc = "", "Версия модели (код: khl_simulation/khl_model, параметры: params.json)"
        val = getattr(params, f.name)
        ws.cell(row=r, column=1, value=f.name)
        ws.cell(row=r, column=2, value=round(val, 4) if isinstance(val, float) else val)
        ws.cell(row=r, column=3, value=status)
        ws.cell(row=r, column=4, value=desc).alignment = Alignment(wrap_text=True)
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BORDER
        r += 1
    ws.conditional_formatting.add(f"C2:C{r}", FormulaRule(formula=['C2="PROVISIONAL"'], fill=PatternFill("solid", fgColor=YELLOW)))
    extra = [
        ("Часовой пояс", "Asia/Yekaterinburg (UTC+5)", "", "Основное время пользователя; в журнале время матча — МСК"),
        ("Источники", "KHL.ru, FONBET, Winline", "", "Протоколы и составы — KHL.ru; линии — обе стороны рынка с отметкой времени"),
        ("Стоп-правило пула", "Банк < 50% стартового", "ПРАВИЛО", "Пул останавливается, модель пересматривается"),
        ("Стоп-правило модели", "CLV ≤ 0 после 150 ставок", "ПРАВИЛО", "Ставки прекращаются до перекалибровки"),
    ]
    for row in extra:
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).border = BORDER
        r += 1


PROMPT = """# khl_simulation — рабочий промпт v2.0 (01.10.2026, Asia/Yekaterinburg UTC+5)
1. Роль: аналитик хоккейных данных. Все вероятности, λ, EV и размеры ставок считает ТОЛЬКО код khl_model (репозиторий, папка khl_simulation). Вручную задавать или «подправлять» λ, вероятности, EV и ставки запрещено. Если код запустить нельзя — в этот день ставок нет.
2. Область запуска: матчи КХЛ ближайших 24 часов, которые ещё не начались. Начавшиеся матчи не анализируются и в журнал не вносятся.
3. Данные о результатах: после каждого игрового дня дописать в results.csv сыгранные матчи — date, season, home, away, home_goals_60, away_goals_60, home_goals_final, away_goals_final, decided (REG/OT/SO). Счёт за 60 минут — из протокола KHL.ru (с голами в пустые ворота). Названия команд — строго как во вкладке «Рейтинги».
4. Данные о матче (matches.csv): match_id, date, time_msk, home, away, home_backup_goalie / away_backup_goalie (1 — подтверждён запасной вратарь), home_b2b / away_b2b (1 — второй матч за 2 дня), home_tz_shift / away_tz_shift (смена часовых поясов за последние 48 ч, часы). Неподтверждённое — оставлять пустым, не угадывать.
5. Коэффициенты (odds.csv): ОБЕ (все) стороны каждого рынка с отметкой времени taken_at: 1X2, DC (1X/X2/12), TOTAL 4.5/5.5/6.5 (O и U), HCP ±1.5 (1 и 2), ML_OT (победитель с ОТ/Б). Без второй стороны маржа снимается грубо.
6. Расчёт: python -m khl_model fit --results results.csv --as-of <дата> --out ratings.json → python -m khl_model predict --matches matches.csv --odds odds.csv --ratings ratings.json --journal journal.csv. Защита от утечки: fit использует только матчи СТРОГО до даты прогноза.
7. Модель v2.0: log λ_home = μ + h + att_home − def_away; log λ_away = μ + att_away − def_home (затухание по времени, регрессия к среднему между сезонами). Матрица счетов за 60 минут считается точно (без Monte Carlo), с поправкой на ничьи (θ) и на голы в пустые ворота. ОТ 3×3 и буллиты — отдельно, сила команд в ОТ сжата к 50/50.
8. Смешивание с рынком: итоговая вероятность = логарифмический пулинг модели и рынка без маржи (вес модели — market_weight_model). EV считается только по итоговой вероятности.
9. Отбор: только строки из picks.csv. Условия: EV ≥ 3% и EV ≥ 0 во всех стресс-сценариях; не больше 1 ставки на матч. Пул определяется коэффициентом: Надёжное < 1.60 ≤ Главное < 2.20 ≤ Смелое. День без ставок — нормальный исход.
10. Банк: у каждого пула свой банк, старт 1 000 у.е. Ставка = ¼ Келли от ТЕКУЩЕГО банка пула, но не больше 3% банка и не меньше 5 у.е. Банк не может уйти в минус. Пул с банком < 50% стартового останавливается.
11. Журнал (вкладка KHL, строки с 8): одна строка на ставку, все поля из picks.csv: Пул, Рынок, Выбор, Линия, Коэффициент, Букмекер, Время кэфа, λ, p модели, q рынка, p итог, p возврата, Fair кэф, EV, EV худший, Kelly, Банк пула до ставки, Ставка. Строку вносить ДО начала матча; после стартового свистка эти поля не меняются.
12. Прогнозы (вкладка «Прогнозы»): одна строка на КАЖДЫЙ разобранный матч, даже без ставки — λ, P1/X/P2 за 60 минут, П1/П2 с ОТ/Б, q1/qX/q2 рынка без маржи. Это основа оценки качества модели.
13. Закрытие линии: за 0–15 минут до начала записать «Кэф закрытия» и «q закрытия» (без маржи по обеим сторонам). CLV = кэф × q закрытия − 1. Устойчивый CLV > 0 — главный признак реального перевеса.
14. Расчёт ставок: после матча внести «Счёт 60м», «Счёт итог» (с пометкой ОТ/Б) и «Результат»: ЗАШЛО / НЕ ЗАШЛО / ВОЗВРАТ / ОТМЕНА; до матча — ОЖИДАНИЕ. P&L, банки, yield, CLV, Brier и цвета считаются формулами и условным форматированием — вручную не вписывать. ML_OT рассчитывается по итоговому счёту, остальные рынки — по счёту за 60 минут.
15. Обоснование: только из признаков модели — разница рейтингов, вратари, b2b, перелёты, расхождение модели и рынка. Сюжетные объяснения без чисел не писать.
16. Качество и калибровка: раз в неделю — python -m khl_model evaluate --journal journal.csv и сравнение Brier/log loss модели и рынка во вкладке «Прогнозы». Раз в месяц (и обязательно после ≥ 300 матчей с начала сезона) — python -m khl_model calibrate --results results.csv; параметры PROVISIONAL во вкладке «Настройки» заменяются откалиброванными. Если после 150 ставок средний CLV ≤ 0 — ставки прекращаются до пересмотра модели.
17. Статистическая честность: по выборке меньше нескольких сотен ставок yield почти ничего не говорит (стандартная ошибка указана в отчёте evaluate). Решения о модели принимаются по CLV и Brier/log loss против рынка, а не по win rate."""


def build_prompt(wb) -> None:
    ws = wb.create_sheet("Промпт")
    ws.column_dimensions["A"].width = 160
    for i, line in enumerate(PROMPT.split("\n"), 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=12)


def copy_archive(wb, src_ws) -> None:
    ws = wb.create_sheet("Журнал v1 (архив)")
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                new = ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new.number_format = cell.number_format
    for r in (1, 7):
        style_header(ws, r, src_ws.max_column if r == 7 else 10)
    for i in range(1, src_ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["AN"].width = 80
    ws.freeze_panes = "A8"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--bets", default=str(ROOT / "data" / "v1_bets.csv"))
    ap.add_argument("--params", default=str(ROOT / "params.json"))
    a = ap.parse_args()
    params = Params.load(a.params)
    bets = pd.read_csv(a.bets)
    src = openpyxl.load_workbook(a.src)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_journal(wb)
    build_forecasts(wb, bets, params)
    totals = build_retro(wb, bets, params)
    build_ratings(wb)
    build_settings(wb, params)
    build_prompt(wb)
    copy_archive(wb, src["KHL"])
    wb.save(a.out)
    print(f"Сохранено: {a.out}; P&L v1 = {totals['v1']:+.0f}, ставок v2 = {totals['v2_bets']}")


if __name__ == "__main__":
    main()
